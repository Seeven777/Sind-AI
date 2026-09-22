import hashlib
import json
import re
import sqlite3
from datetime import datetime
from pathlib import Path
import urllib.parse

import requests
from bs4 import BeautifulSoup


class KnowledgeBase:
    def __init__(self, db_path, storage_dir, user_agent="JarvisSindPet/0.7"):
        self.db_path = Path(db_path)
        self.storage_dir = Path(storage_dir)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        self.user_agent = user_agent
        self._init_db()

    def _connect(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self):
        with self._connect() as conn:
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS kb_documents (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    collection TEXT NOT NULL,
                    title TEXT NOT NULL,
                    source TEXT NOT NULL,
                    source_type TEXT NOT NULL,
                    source_hash TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    metadata_json TEXT NOT NULL DEFAULT '{}'
                );

                CREATE TABLE IF NOT EXISTS kb_chunks (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    document_id INTEGER NOT NULL,
                    collection TEXT NOT NULL,
                    chunk_index INTEGER NOT NULL,
                    text TEXT NOT NULL,
                    FOREIGN KEY(document_id) REFERENCES kb_documents(id)
                );

                CREATE INDEX IF NOT EXISTS idx_kb_docs_collection
                ON kb_documents(collection);

                CREATE INDEX IF NOT EXISTS idx_kb_chunks_collection
                ON kb_chunks(collection);
            """)
            try:
                conn.execute("""
                    CREATE VIRTUAL TABLE IF NOT EXISTS kb_fts
                    USING fts5(text, collection UNINDEXED, document_id UNINDEXED, chunk_id UNINDEXED);
                """)
            except sqlite3.OperationalError:
                pass

    def _slug_collection(self, value):
        value = re.sub(r"\s+", "_", str(value).strip().lower())
        value = re.sub(r"[^\w.-]+", "", value)
        return value or "default"

    def _chunk(self, text, chunk_chars=1400, overlap=180):
        text = re.sub(r"\r\n?", "\n", str(text))
        paragraphs = [re.sub(r"\s+", " ", p).strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        chunks = []
        current = ""

        for p in paragraphs:
            if len(current) + len(p) + 2 <= chunk_chars:
                current = (current + "\n\n" + p).strip()
                continue

            if current:
                chunks.append(current)

            if len(p) <= chunk_chars:
                current = p
            else:
                start = 0
                while start < len(p):
                    end = min(len(p), start + chunk_chars)
                    chunks.append(p[start:end])
                    if end == len(p):
                        current = ""
                        break
                    start = max(end - overlap, start + 1)

        if current:
            chunks.append(current)

        return chunks or [text[:chunk_chars]]

    def _delete_document_id(self, conn, doc_id):
        rows = conn.execute("SELECT id FROM kb_chunks WHERE document_id=?", (doc_id,)).fetchall()
        chunk_ids = [int(r["id"]) for r in rows]
        if chunk_ids:
            placeholders = ",".join("?" for _ in chunk_ids)
            try:
                conn.execute(f"DELETE FROM kb_fts WHERE chunk_id IN ({placeholders})", chunk_ids)
            except sqlite3.OperationalError:
                pass
        conn.execute("DELETE FROM kb_chunks WHERE document_id=?", (doc_id,))
        conn.execute("DELETE FROM kb_documents WHERE id=?", (doc_id,))

    def ingest_text(self, collection, title, text, source="manual", source_type="text", metadata=None):
        collection = self._slug_collection(collection)
        title = str(title).strip() or "Documento"
        text = str(text).strip()
        if not text:
            return {"ok": False, "error": "Conteúdo vazio."}

        source_hash = hashlib.sha256((str(source) + "\n" + text).encode("utf-8")).hexdigest()
        now = datetime.now().isoformat(timespec="seconds")

        with self._connect() as conn:
            existing = conn.execute(
                "SELECT id FROM kb_documents WHERE collection=? AND source_hash=?",
                (collection, source_hash)
            ).fetchone()
            if existing:
                return {"ok": True, "document_id": int(existing["id"]), "collection": collection, "duplicate": True}

            cur = conn.execute(
                """INSERT INTO kb_documents(collection,title,source,source_type,source_hash,created_at,metadata_json)
                   VALUES(?,?,?,?,?,?,?)""",
                (collection, title, str(source), source_type, source_hash, now, json.dumps(metadata or {}, ensure_ascii=False))
            )
            doc_id = int(cur.lastrowid)

            chunks = self._chunk(text)
            for idx, chunk in enumerate(chunks):
                ccur = conn.execute(
                    "INSERT INTO kb_chunks(document_id,collection,chunk_index,text) VALUES(?,?,?,?)",
                    (doc_id, collection, idx, chunk)
                )
                chunk_id = int(ccur.lastrowid)
                try:
                    conn.execute(
                        "INSERT INTO kb_fts(text,collection,document_id,chunk_id) VALUES(?,?,?,?)",
                        (chunk, collection, doc_id, chunk_id)
                    )
                except sqlite3.OperationalError:
                    pass

        return {
            "ok": True,
            "document_id": doc_id,
            "collection": collection,
            "chunks": len(chunks),
            "chars": len(text),
        }

    def _extract_file(self, path):
        p = Path(path).expanduser().resolve()
        if not p.is_file():
            raise FileNotFoundError(str(p))
        ext = p.suffix.lower()

        if ext in {".txt", ".md", ".csv", ".json", ".log", ".html", ".htm"}:
            raw = p.read_text(encoding="utf-8", errors="replace")
            if ext in {".html", ".htm"}:
                raw = BeautifulSoup(raw, "html.parser").get_text("\n", strip=True)
            return raw, ext.lstrip(".")

        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
            return text, "pdf"

        if ext == ".docx":
            from docx import Document
            doc = Document(str(p))
            text = "\n".join(par.text for par in doc.paragraphs)
            return text, "docx"

        raise ValueError(f"Tipo de arquivo não suportado: {ext}")

    def ingest_file(self, collection, path, title=None):
        try:
            p = Path(path).expanduser().resolve()
            text, kind = self._extract_file(p)
            return self.ingest_text(
                collection=collection,
                title=title or p.name,
                text=text,
                source=str(p),
                source_type=kind,
                metadata={"path": str(p)},
            )
        except Exception as exc:
            return {"ok": False, "error": str(exc)}

    def ingest_url(self, collection, url, title=None):
        try:
            url = str(url).strip()
            if not url.startswith("https://"):
                return {"ok": False, "error": "Somente URLs HTTPS são permitidas."}
            r = requests.get(url, timeout=20, headers={"User-Agent": self.user_agent})
            r.raise_for_status()
            ctype = (r.headers.get("content-type") or "").lower()

            if "html" in ctype:
                soup = BeautifulSoup(r.text, "html.parser")
                for tag in soup(["script", "style", "noscript"]):
                    tag.decompose()
                text = soup.get_text("\n", strip=True)
                page_title = soup.title.get_text(" ", strip=True) if soup.title else url
            else:
                text = r.text
                page_title = url

            return self.ingest_text(
                collection=collection,
                title=title or page_title,
                text=text,
                source=url,
                source_type="url",
                metadata={"url": url, "content_type": ctype},
            )
        except Exception as exc:
            return {"ok": False, "error": str(exc)}

    def search(self, collection, query, limit=8):
        collection = self._slug_collection(collection)
        query = str(query).strip()
        if not query:
            return {"ok": True, "items": [], "count": 0}

        with self._connect() as conn:
            items = []
            try:
                safe = " ".join(re.findall(r"[\wÀ-ÿ.-]+", query))
                rows = conn.execute(
                    """SELECT f.text, f.document_id, f.chunk_id, d.title, d.source,
                              bm25(kb_fts) AS score
                       FROM kb_fts f
                       JOIN kb_documents d ON d.id=f.document_id
                       WHERE kb_fts MATCH ? AND f.collection=?
                       ORDER BY score
                       LIMIT ?""",
                    (safe, collection, int(limit))
                ).fetchall()
                for r in rows:
                    items.append({
                        "document_id": r["document_id"],
                        "chunk_id": r["chunk_id"],
                        "title": r["title"],
                        "source": r["source"],
                        "text": r["text"],
                        "score": r["score"],
                    })
            except sqlite3.OperationalError:
                like = f"%{query}%"
                rows = conn.execute(
                    """SELECT c.text,c.document_id,c.id AS chunk_id,d.title,d.source
                       FROM kb_chunks c JOIN kb_documents d ON d.id=c.document_id
                       WHERE c.collection=? AND c.text LIKE ?
                       LIMIT ?""",
                    (collection, like, int(limit))
                ).fetchall()
                items = [dict(r) for r in rows]

        return {"ok": True, "collection": collection, "query": query, "items": items, "count": len(items)}

    def list_collections(self):
        with self._connect() as conn:
            rows = conn.execute(
                """SELECT collection, COUNT(*) docs
                   FROM kb_documents GROUP BY collection ORDER BY collection"""
            ).fetchall()
        return {"ok": True, "items": [dict(r) for r in rows], "count": len(rows)}

    def list_documents(self, collection, limit=100):
        collection = self._slug_collection(collection)
        with self._connect() as conn:
            rows = conn.execute(
                """SELECT id,collection,title,source,source_type,created_at,metadata_json
                   FROM kb_documents WHERE collection=? ORDER BY id DESC LIMIT ?""",
                (collection, int(limit))
            ).fetchall()
        items = []
        for r in rows:
            d = dict(r)
            try:
                d["metadata"] = json.loads(d.pop("metadata_json"))
            except Exception:
                pass
            items.append(d)
        return {"ok": True, "items": items, "count": len(items)}

    def document_info(self, document_id):
        with self._connect() as conn:
            row = conn.execute("SELECT * FROM kb_documents WHERE id=?", (int(document_id),)).fetchone()
            if not row:
                return {"ok": False, "error": "Documento não encontrado."}
            chunks = conn.execute("SELECT COUNT(*) n FROM kb_chunks WHERE document_id=?", (int(document_id),)).fetchone()["n"]
        data = dict(row)
        data["chunks"] = chunks
        return {"ok": True, "data": data}

    def remove_document(self, document_id):
        with self._connect() as conn:
            row = conn.execute("SELECT id,title FROM kb_documents WHERE id=?", (int(document_id),)).fetchone()
            if not row:
                return {"ok": False, "error": "Documento não encontrado."}
            self._delete_document_id(conn, int(document_id))
        return {"ok": True, "document_id": int(document_id), "title": row["title"]}

    def clear_collection(self, collection):
        collection = self._slug_collection(collection)
        with self._connect() as conn:
            rows = conn.execute("SELECT id FROM kb_documents WHERE collection=?", (collection,)).fetchall()
            for r in rows:
                self._delete_document_id(conn, int(r["id"]))
        return {"ok": True, "collection": collection, "removed": len(rows)}

    def stats(self):
        with self._connect() as conn:
            docs = conn.execute("SELECT COUNT(*) n FROM kb_documents").fetchone()["n"]
            chunks = conn.execute("SELECT COUNT(*) n FROM kb_chunks").fetchone()["n"]
            cols = conn.execute("SELECT COUNT(DISTINCT collection) n FROM kb_documents").fetchone()["n"]
        return {"ok": True, "collections": cols, "documents": docs, "chunks": chunks}

    def export_collection(self, collection, path):
        collection = self._slug_collection(collection)
        out = Path(path).expanduser().resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        docs = self.list_documents(collection, limit=10000)["items"]
        with self._connect() as conn:
            payload = []
            for d in docs:
                chunks = conn.execute(
                    "SELECT chunk_index,text FROM kb_chunks WHERE document_id=? ORDER BY chunk_index",
                    (d["id"],)
                ).fetchall()
                payload.append({
                    "document": d,
                    "chunks": [dict(c) for c in chunks],
                })
        out.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
        return {"ok": True, "path": str(out), "documents": len(payload)}
