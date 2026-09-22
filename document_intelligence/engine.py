import csv
import difflib
import hashlib
import json
import mimetypes
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup


SUPPORTED = {'.txt','.md','.log','.csv','.json','.html','.htm','.pdf','.docx'}


class DocumentIntelligence:
    def __init__(self, allowed_root=None):
        self.allowed_root = Path(allowed_root).resolve() if allowed_root else None

    def _path(self, value):
        p=Path(value).expanduser().resolve()
        if self.allowed_root and self.allowed_root not in p.parents and p != self.allowed_root:
            raise PermissionError('O arquivo está fora da área permitida.')
        return p

    def _hash(self,p,algo='sha256'):
        h=hashlib.new(algo)
        with p.open('rb') as fh:
            for chunk in iter(lambda:fh.read(1024*1024),b''): h.update(chunk)
        return h.hexdigest()

    def _extract(self,p):
        ext=p.suffix.lower()
        if ext in {'.txt','.md','.log','.csv','.json'}:
            return p.read_text(encoding='utf-8',errors='replace'), ext.lstrip('.')
        if ext in {'.html','.htm'}:
            raw=p.read_text(encoding='utf-8',errors='replace')
            soup=BeautifulSoup(raw,'html.parser')
            for t in soup(['script','style','noscript']): t.decompose()
            return soup.get_text('\n',strip=True),'html'
        if ext=='.pdf':
            from pypdf import PdfReader
            r=PdfReader(str(p)); txt='\n\n'.join((pg.extract_text() or '') for pg in r.pages)
            return txt,'pdf'
        if ext=='.docx':
            from docx import Document
            d=Document(str(p)); txt='\n'.join(x.text for x in d.paragraphs)
            return txt,'docx'
        raise ValueError(f'Tipo não suportado: {ext}')

    def inspect(self,path):
        try:
            p=self._path(path)
            if not p.exists(): return {'ok':False,'error':'Arquivo não encontrado.'}
            st=p.stat()
            return {'ok':True,'path':str(p),'name':p.name,'extension':p.suffix.lower(),'size':st.st_size,
                    'modified':datetime.fromtimestamp(st.st_mtime).isoformat(timespec='seconds'),
                    'mime':mimetypes.guess_type(str(p))[0] or '', 'supported':p.suffix.lower() in SUPPORTED,
                    'is_file':p.is_file(),'is_dir':p.is_dir()}
        except Exception as exc: return {'ok':False,'error':str(exc)}

    def extract_text(self,path,max_chars=30000):
        try:
            p=self._path(path); text,kind=self._extract(p)
            return {'ok':True,'path':str(p),'kind':kind,'chars':len(text),'text':text[:int(max_chars)]}
        except Exception as exc: return {'ok':False,'error':str(exc)}

    def pdf_info(self,path):
        try:
            from pypdf import PdfReader
            p=self._path(path); r=PdfReader(str(p)); meta={str(k):str(v) for k,v in (r.metadata or {}).items()}
            return {'ok':True,'pages':len(r.pages),'metadata':meta,'path':str(p)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def pdf_page(self,path,page=1,max_chars=12000):
        try:
            from pypdf import PdfReader
            p=self._path(path); r=PdfReader(str(p)); idx=int(page)-1
            if idx<0 or idx>=len(r.pages): return {'ok':False,'error':'Página inválida.'}
            text=r.pages[idx].extract_text() or ''
            return {'ok':True,'page':idx+1,'text':text[:int(max_chars)]}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def pdf_search(self,path,query,regex=False,case_sensitive=False,limit=60):
        try:
            from pypdf import PdfReader
            p=self._path(path); r=PdfReader(str(p)); items=[]
            flags=0 if case_sensitive else re.I
            pattern=re.compile(query,flags) if regex else None
            needle=str(query) if case_sensitive else str(query).lower()
            for i,pg in enumerate(r.pages,1):
                txt=pg.extract_text() or ''
                hay=txt if case_sensitive else txt.lower()
                matched=bool(pattern.search(txt)) if pattern else needle in hay
                if matched:
                    pos=(pattern.search(txt).start() if pattern and pattern.search(txt) else hay.find(needle))
                    excerpt=re.sub(r'\s+',' ',txt[max(0,pos-180):pos+420]).strip()
                    items.append({'page':i,'excerpt':excerpt})
                    if len(items)>=int(limit):break
            return {'ok':True,'query':query,'items':items,'count':len(items)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def docx_info(self,path):
        try:
            from docx import Document
            p=self._path(path); d=Document(str(p));
            headings=[{'level':par.style.name,'text':par.text} for par in d.paragraphs if par.text.strip() and par.style and par.style.name.lower().startswith('heading')]
            return {'ok':True,'paragraphs':len(d.paragraphs),'tables':len(d.tables),'headings':headings[:100]}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def docx_tables(self,path,limit=30):
        try:
            from docx import Document
            p=self._path(path); d=Document(str(p)); items=[]
            for ti,t in enumerate(d.tables[:int(limit)],1):
                rows=[[c.text for c in row.cells] for row in t.rows]
                items.append({'table':ti,'rows':rows})
            return {'ok':True,'items':items,'count':len(items)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def headings(self,path):
        try:
            p=self._path(path); ext=p.suffix.lower()
            if ext=='.docx': return self.docx_info(path)
            text,_=self._extract(p); items=[]
            for line in text.splitlines():
                s=line.strip()
                if not s: continue
                if re.match(r'^#{1,6}\s+',s): items.append({'text':re.sub(r'^#+\s*','',s),'kind':'markdown'})
                elif len(s)<=120 and (s.isupper() or re.match(r'^\d+(?:\.\d+)*\s+\S+',s)):
                    items.append({'text':s,'kind':'heuristic'})
            return {'ok':True,'items':items[:200],'count':len(items)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def search_text(self,path,query,regex=False,limit=100):
        try:
            p=self._path(path); text,_=self._extract(p); lines=text.splitlines(); items=[]
            patt=re.compile(query,re.I) if regex else None; needle=str(query).lower()
            for i,line in enumerate(lines,1):
                hit=bool(patt.search(line)) if patt else needle in line.lower()
                if hit:
                    items.append({'line':i,'text':line[:800]})
                    if len(items)>=int(limit):break
            return {'ok':True,'items':items,'count':len(items)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def metrics(self,path):
        r=self.extract_text(path,max_chars=2_000_000)
        if not r.get('ok'):return r
        text=r['text']; words=re.findall(r'\b[\wÀ-ÿ-]+\b',text)
        return {'ok':True,'chars':len(text),'words':len(words),'lines':len(text.splitlines()),
                'paragraphs':len([p for p in re.split(r'\n\s*\n',text) if p.strip()]),
                'reading_minutes':round(len(words)/200,2)}

    def top_terms(self,path,limit=30):
        r=self.extract_text(path,max_chars=2_000_000)
        if not r.get('ok'):return r
        stop={'para','com','uma','que','dos','das','por','de','do','da','em','no','na','e','o','a','os','as','um'}
        words=[w.lower() for w in re.findall(r'\b[\wÀ-ÿ-]{3,}\b',r['text']) if w.lower() not in stop]
        return {'ok':True,'items':[{'term':w,'count':c} for w,c in Counter(words).most_common(int(limit))]}

    def compare(self,path_a,path_b):
        a=self.extract_text(path_a,max_chars=2_000_000); b=self.extract_text(path_b,max_chars=2_000_000)
        if not a.get('ok'):return a
        if not b.get('ok'):return b
        ratio=difflib.SequenceMatcher(None,a['text'],b['text']).ratio()
        diff='\n'.join(difflib.unified_diff(a['text'].splitlines(),b['text'].splitlines(),fromfile=str(path_a),tofile=str(path_b),lineterm=''))
        return {'ok':True,'similarity':round(ratio,4),'diff':diff[:30000]}

    def inventory(self,folder,recursive=True,limit=2000):
        try:
            root=self._path(folder)
            it=root.rglob('*') if recursive else root.glob('*'); items=[]
            for p in it:
                if not p.is_file():continue
                st=p.stat(); items.append({'path':str(p),'name':p.name,'extension':p.suffix.lower(),'size':st.st_size,'modified':st.st_mtime,'supported':p.suffix.lower() in SUPPORTED})
                if len(items)>=int(limit):break
            return {'ok':True,'root':str(root),'items':items,'count':len(items)}
        except Exception as exc:return {'ok':False,'error':str(exc)}

    def folder_summary(self,folder):
        r=self.inventory(folder,True,5000)
        if not r.get('ok'):return r
        ext=Counter(x['extension'] or '(sem extensão)' for x in r['items']); total=sum(x['size'] for x in r['items'])
        supported=sum(1 for x in r['items'] if x['supported'])
        return {'ok':True,'files':len(r['items']),'bytes':total,'supported_documents':supported,'extensions':dict(ext.most_common())}

    def duplicates(self,folder,limit=200):
        r=self.inventory(folder,True,5000)
        if not r.get('ok'):return r
        groups=defaultdict(list)
        size_groups=defaultdict(list)
        for item in r['items']:
            if item['size']>0:size_groups[item['size']].append(item)
        for size,items in size_groups.items():
            if len(items)<2:continue
            for item in items:
                try: groups[self._hash(Path(item['path']))].append(item['path'])
                except Exception:pass
        dup=[{'sha256':h,'paths':paths,'count':len(paths)} for h,paths in groups.items() if len(paths)>1][:int(limit)]
        return {'ok':True,'items':dup,'count':len(dup)}

    def recent_files(self,folder,limit=30):
        r=self.inventory(folder,True,5000)
        if not r.get('ok'):return r
        items=sorted(r['items'],key=lambda x:-x['modified'])[:int(limit)]
        for x in items:x['modified_iso']=datetime.fromtimestamp(x['modified']).isoformat(timespec='seconds')
        return {'ok':True,'items':items,'count':len(items)}

    def large_files(self,folder,limit=30):
        r=self.inventory(folder,True,5000)
        if not r.get('ok'):return r
        items=sorted(r['items'],key=lambda x:-x['size'])[:int(limit)]
        return {'ok':True,'items':items,'count':len(items)}

    def batch_search(self,folder,query,limit=100):
        r=self.inventory(folder,True,1200)
        if not r.get('ok'):return r
        results=[]
        for item in r['items']:
            if not item['supported']:continue
            s=self.search_text(item['path'],query,False,5)
            if s.get('count'):
                results.append({'path':item['path'],'matches':s['items']})
                if len(results)>=int(limit):break
        return {'ok':True,'items':results,'count':len(results)}

    def manifest(self,folder):
        r=self.inventory(folder,True,5000)
        if not r.get('ok'):return r
        items=[]
        for x in r['items']:
            try: sha=self._hash(Path(x['path']))
            except Exception: sha=''
            items.append({**x,'sha256':sha})
        return {'ok':True,'root':r['root'],'items':items,'count':len(items)}

    def execute(self,operation,**p):
        path=p.get('path',''); folder=p.get('folder','')
        if operation=='inspect':return self.inspect(path)
        if operation=='sha256':
            try:q=self._path(path);return {'ok':True,'sha256':self._hash(q)}
            except Exception as exc:return {'ok':False,'error':str(exc)}
        if operation=='md5':
            try:q=self._path(path);return {'ok':True,'md5':self._hash(q,'md5')}
            except Exception as exc:return {'ok':False,'error':str(exc)}
        if operation=='extract_text':return self.extract_text(path,p.get('max_chars',30000))
        if operation=='metrics':return self.metrics(path)
        if operation=='top_terms':return self.top_terms(path,p.get('limit',30))
        if operation=='headings':return self.headings(path)
        if operation=='search':return self.search_text(path,p.get('query',''),p.get('regex',False),p.get('limit',100))
        if operation=='pdf_info':return self.pdf_info(path)
        if operation=='pdf_page':return self.pdf_page(path,p.get('page',1),p.get('max_chars',12000))
        if operation=='pdf_search':return self.pdf_search(path,p.get('query',''),p.get('regex',False),p.get('case_sensitive',False),p.get('limit',60))
        if operation=='docx_info':return self.docx_info(path)
        if operation=='docx_tables':return self.docx_tables(path,p.get('limit',30))
        if operation=='compare':return self.compare(p.get('path_a',''),p.get('path_b',''))
        if operation=='inventory':return self.inventory(folder,p.get('recursive',True),p.get('limit',2000))
        if operation=='folder_summary':return self.folder_summary(folder)
        if operation=='duplicates':return self.duplicates(folder,p.get('limit',200))
        if operation=='recent_files':return self.recent_files(folder,p.get('limit',30))
        if operation=='large_files':return self.large_files(folder,p.get('limit',30))
        if operation=='batch_search':return self.batch_search(folder,p.get('query',''),p.get('limit',100))
        if operation=='manifest':return self.manifest(folder)
        return {'ok':False,'error':f'Operação documental desconhecida: {operation}'}
